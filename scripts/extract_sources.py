# -*- coding: utf-8 -*-
"""扫描课程目录，抽取课件 PPT / AI纪要 docx / 作业 PDF 文本，并判定主导语言。

用法:
    python extract_sources.py <课程目录> [--out <输出目录>]

输出:
    <out>/course_sources.md   全部文本素材（含盘点清单）
    <out>/language.txt        判定结果: en 或 zh
"""
import argparse
import re
import sys
from pathlib import Path

CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")
ASCII_LETTER_RE = re.compile(r"[A-Za-z]")

AUDIO_EXTS = {".m4a", ".mp3", ".wav", ".aac", ".flac", ".ogg", ".wma"}

def pdf_text(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        sys.stderr.write("缺少 PyMuPDF，请先: pip install -r requirements.txt\n")
        raise
    doc = fitz.open(str(path))
    parts = []
    for page in doc:
        parts.append(page.get_text())
    doc.close()
    return "\n".join(parts)

def docx_text(path: Path) -> str:
    try:
        import docx
    except ImportError:
        sys.stderr.write("缺少 python-docx，请先: pip install -r requirements.txt\n")
        raise
    d = docx.Document(str(path))
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)

def classify_file(path: Path):
    name = path.name
    low = name.lower()
    parent = path.parent.name.lower()
    ext = path.suffix.lower()
    if ext in AUDIO_EXTS:
        return "audio"
    # 纪要 docx 优先判定
    if ext == ".docx":
        return "minutes"
    if ext != ".pdf":
        return None
    if ("AI纪要" in name) or ("纪要" in name):
        # 兼容 PDF 形式的纪要
        return "minutes"
    if "ppt" in parent or "slide" in low or "lecture" in low:
        return "ppt"
    if "homework" in parent or "assignment" in low or low.startswith("hw") or "作业" in name:
        return "homework"
    if re.match(r"chapter\s*\d", low) or re.match(r"^\d+\.\d+", low):
        return "ppt"
    # 既有生成物（复习/quiz/知识体系等）不作为输入，跳过
    if re.search(r"review|quiz|guide|knowledge|summary", low):
        return None
    return "ppt"  # 根目录下无法判定的 PDF 默认视作课件

def scan(course_dir: Path):
    found = {"ppt": [], "minutes": [], "homework": [], "audio": []}
    for p in sorted(course_dir.rglob("*")):
        if not p.is_file():
            continue
        if ".git" in p.parts:
            continue
        kind = classify_file(p)
        if kind:
            found[kind].append(p)
    return found

def detect_language(ppt_text: str) -> str:
    zh = len(CJK_RE.findall(ppt_text))
    en = len(ASCII_LETTER_RE.findall(ppt_text))
    if zh + en == 0:
        return "en"
    return "zh" if zh / (zh + en) >= 0.5 else "en"

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("course_dir", type=str)
    ap.add_argument("--out", type=str, default=None)
    args = ap.parse_args()

    course = Path(args.course_dir).resolve()
    if not course.is_dir():
        sys.stderr.write(f"课程目录不存在: {course}\n")
        sys.exit(2)
    out = Path(args.out).resolve() if args.out else course / "review"
    out.mkdir(parents=True, exist_ok=True)

    found = scan(course)
    lines = []
    lines.append(f"# 课程素材盘点 Course Sources: {course.name}\n")
    lines.append(f"- 课件 PPT/PDF: {len(found['ppt'])} 个")
    lines.append(f"- AI纪要 docx: {len(found['minutes'])} 个")
    lines.append(f"- 作业 PDF: {len(found['homework'])} 个")
    lines.append(f"- 音频文件: {len(found['audio'])} 个\n")

    # 语言判定只看 PPT 文本
    ppt_text_all = ""
    for f in found["ppt"]:
        ppt_text_all += pdf_text(f) + "\n\n"
    lang = detect_language(ppt_text_all)
    lines.append(f"## 判定语言 Language: **{lang}**\n")
    lines.append("> 若为 `en`：英文为主 + 陌生术语中文注释；若为 `zh`：中文为主。详见 references/language-rule.md\n")

    sections = [("课件 PPT", found["ppt"], pdf_text),
                ("AI纪要 Minutes", found["minutes"], docx_text),
                ("作业 Homework", found["homework"], pdf_text)]
    for title, files, extractor in sections:
        lines.append(f"\n---\n\n# {title}（{len(files)}）\n")
        if not files:
            lines.append("_（无）_\n")
            continue
        for f in files:
            rel = f.relative_to(course)
            lines.append(f"\n## {rel}\n")
            try:
                text = extractor(f)
            except Exception as e:
                lines.append(f"\n> 读取失败: {e}\n")
                continue
            text = text.strip()
            lines.append(text + "\n")

    src_path = out / "course_sources.md"
    src_path.write_text("\n".join(lines), encoding="utf-8")
    (out / "language.txt").write_text(lang + "\n", encoding="utf-8")

    print(f"language={lang}")
    print(f"ppt={len(found['ppt'])} minutes={len(found['minutes'])} homework={len(found['homework'])} audio={len(found['audio'])}")
    print(f"wrote {src_path}")
    print(f"wrote {out / 'language.txt'}")

if __name__ == "__main__":
    main()
